"""Unit tests for the resume parser modular package (app/services/resume_parser/)."""

import io
import json
import pytest
import docx
import pypdf

from app.services.resume_parser.clean import (
    clean_text,
    normalize_bullets,
    normalize_whitespace,
    remove_control_characters,
    remove_headers_and_footers,
)
from app.services.resume_parser.extract import (
    CorruptedFileError,
    EmptyFileError,
    UnsupportedFileTypeError,
    extract_from_docx,
    extract_from_pdf,
    extract_from_txt,
    extract_text,
)
from app.services.resume_parser.parser import (
    ResumeParsingError,
    parse_resume,
)
from app.services.resume_parser.schema import (
    EducationItem,
    ExperienceItem,
    ProjectItem,
    StructuredResume,
)
from app.services.resume_parser.structure import (
    build_structuring_prompt,
    heuristic_structure_resume,
    parse_llm_response,
    structure_resume,
)


# =============================================================================
# Helper Utilities for Generating In-Memory Test Files
# =============================================================================

def create_sample_pdf_bytes(text_lines: list[str]) -> bytes:
    """Generate a valid in-memory PDF using pypdf."""
    writer = pypdf.PdfWriter()
    # Add a blank page
    page = writer.add_blank_page(width=612, height=792)
    # pypdf does not have a rich font drawing engine built-in, but we can test
    # extraction or write a simple stream or mock the reader if needed.
    # Alternatively, write a minimal PDF object stream:
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def create_sample_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] = None) -> bytes:
    """Generate a valid in-memory DOCX using python-docx."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell in enumerate(row):
                table.cell(r_idx, c_idx).text = cell
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


SAMPLE_RESUME_TEXT = """
Jane Doe
jane.doe@example.com
(555) 123-4567

Summary:
Experienced Software Engineer with a passion for building robust distributed systems.

Skills:
Python, FastAPI, Docker, PostgreSQL, React, AWS

Experience:
Senior Software Engineer at Tech Corp
- Built real-time streaming backend handling 50k req/s
- Led migration of legacy monolith to microservices architecture

Software Engineer - Startup Inc
- Developed REST APIs and background task workers

Education:
B.S. in Computer Science from State University
- Graduated magna cum laude

Projects:
MockWise
- AI mock interview platform utilizing speech-to-speech pipelines
- Integrated automated resume parsing
"""


# =============================================================================
# 1. Schema Tests (schema.py)
# =============================================================================

def test_schema_instantiation_and_to_db_json():
    """Verify StructuredResume model creation and to_db_json serialization."""
    resume = StructuredResume(
        candidate_name="Jane Doe",
        email="jane.doe@example.com",
        phone="555-123-4567",
        skills=["Python", "FastAPI"],
        experience=[
            ExperienceItem(
                role="Software Engineer",
                company="Tech Corp",
                highlights=["Built APIs"],
            )
        ],
        education=[
            EducationItem(
                degree="B.S. Computer Science",
                institution="MIT",
            )
        ],
        projects=[
            ProjectItem(
                title="MockWise",
                technologies=["FastAPI", "SQLite"],
            )
        ],
        raw_text="Sample raw resume content",
    )

    db_data = resume.to_db_json()
    assert "skills_json" in db_data
    assert "projects_json" in db_data
    assert "experience_json" in db_data
    assert "education_json" in db_data
    assert "raw_text" in db_data

    # Check that serialized fields parse back properly
    skills = json.loads(db_data["skills_json"])
    assert skills == ["Python", "FastAPI"]

    projects = json.loads(db_data["projects_json"])
    assert len(projects) == 1
    assert projects[0]["title"] == "MockWise"

    experience = json.loads(db_data["experience_json"])
    assert len(experience) == 1
    assert experience[0]["company"] == "Tech Corp"

    education = json.loads(db_data["education_json"])
    assert len(education) == 1
    assert education[0]["degree"] == "B.S. Computer Science"

    assert db_data["raw_text"] == "Sample raw resume content"


# =============================================================================
# 2. Cleaning Tests (clean.py)
# =============================================================================

def test_remove_control_characters():
    """Verify non-printable control characters are stripped while preserving whitespace."""
    dirty = "Hello\x00World\x08!\nTab\tseparated\r\n"
    cleaned = remove_control_characters(dirty)
    assert cleaned == "HelloWorld!\nTab\tseparated\r\n"


def test_remove_headers_and_footers():
    """Verify page number headers and footers are removed."""
    raw = "Header content\nPage 1 of 3\nResume body\n2 / 3\nFooter"
    cleaned = remove_headers_and_footers(raw)
    assert "Page 1 of 3" not in cleaned
    assert "2 / 3" not in cleaned
    assert "Resume body" in cleaned


def test_normalize_bullets():
    """Verify various unicode bullet points normalize to markdown dashes."""
    text = "Skills:\n• Python\n· Docker\n▪ FastAPI\n* AWS"
    normalized = normalize_bullets(text)
    assert "- Python" in normalized
    assert "- Docker" in normalized
    assert "- FastAPI" in normalized
    assert "- AWS" in normalized


def test_normalize_whitespace():
    """Verify excessive spaces and multiple blank lines are collapsed."""
    messy = "  Line   with    lots    of   spaces  \n\n\n\n\nNext line  "
    cleaned = normalize_whitespace(messy)
    assert cleaned == "Line with lots of spaces\n\nNext line"


def test_clean_text_pipeline():
    """Verify full cleaning pipeline with empty, None, and complex inputs."""
    assert clean_text("") == ""
    assert clean_text(None) == ""
    assert clean_text("   \n\n\t  ") == ""

    dirty_resume = "• Skill A\x00\nPage 1 of 2\n\n\n• Skill B"
    result = clean_text(dirty_resume)
    assert "\x00" not in result
    assert "Page 1 of 2" not in result
    assert "- Skill A" in result
    assert "- Skill B" in result


# =============================================================================
# 3. Extraction Tests (extract.py)
# =============================================================================

def test_extract_from_txt():
    """Verify plain text extraction from UTF-8 and Latin-1."""
    content = "Candidate Name: Alex\nSkills: Python, SQL".encode("utf-8")
    extracted = extract_from_txt(content)
    assert "Candidate Name: Alex" in extracted

    # Latin-1 encoded bytes
    latin1_content = "Résumé of José".encode("latin-1")
    extracted_latin1 = extract_from_txt(latin1_content)
    assert "José" in extracted_latin1


def test_extract_from_pdf(monkeypatch):
    """Verify PDF extraction across multiple pages."""
    class MockPage:
        def __init__(self, text):
            self._text = text
        def extract_text(self):
            return self._text

    class MockReader:
        def __init__(self, stream):
            self.pages = [MockPage("Candidate: Jane"), MockPage("Experience: 5 years")]

    monkeypatch.setattr("pypdf.PdfReader", MockReader)
    extracted = extract_from_pdf(b"fake pdf stream")
    assert "Candidate: Jane" in extracted
    assert "Experience: 5 years" in extracted


def test_extract_from_docx():
    """Verify extraction from DOCX paragraphs and tables."""
    docx_bytes = create_sample_docx_bytes(
        paragraphs=["Alex Smith", "Software Engineer"],
        table_rows=[["Skill", "Level"], ["Python", "Expert"]],
    )
    extracted = extract_from_docx(docx_bytes)
    assert "Alex Smith" in extracted
    assert "Software Engineer" in extracted
    assert "Python | Expert" in extracted


def test_extract_empty_file_raises():
    """Verify EmptyFileError is raised on 0-byte file."""
    with pytest.raises(EmptyFileError):
        extract_text(b"", "resume.txt")


def test_extract_unsupported_extension_raises():
    """Verify UnsupportedFileTypeError is raised for unknown file formats."""
    with pytest.raises(UnsupportedFileTypeError) as exc:
        extract_text(b"some content", "resume.exe")
    assert "Unsupported file format" in str(exc.value)


def test_extract_corrupted_file_raises():
    """Verify CorruptedFileError is raised when file bytes are invalid."""
    with pytest.raises(CorruptedFileError):
        extract_text(b"not a valid pdf", "resume.pdf")

    with pytest.raises(CorruptedFileError):
        extract_text(b"not a valid docx", "resume.docx")


# =============================================================================
# 4. Structuring Tests (structure.py)
# =============================================================================

def test_heuristic_structuring():
    """Verify heuristic parser parses candidate name, email, phone, and sections."""
    parsed = heuristic_structure_resume(SAMPLE_RESUME_TEXT)
    assert parsed.candidate_name == "Jane Doe"
    assert parsed.email == "jane.doe@example.com"
    assert parsed.phone == "(555) 123-4567"
    assert "Python" in parsed.skills
    assert "FastAPI" in parsed.skills
    assert len(parsed.experience) >= 1
    assert len(parsed.education) >= 1
    assert len(parsed.projects) >= 1


def test_structure_resume_with_llm_callable():
    """Verify structure_resume invokes and parses LLM response correctly."""
    mock_llm_json = json.dumps({
        "candidate_name": "Bob Builder",
        "email": "bob@builder.com",
        "phone": "123-456-7890",
        "summary": "Full Stack Architect",
        "skills": ["Go", "Kubernetes"],
        "experience": [
            {
                "role": "Lead Architect",
                "company": "Constructions Inc",
                "highlights": ["Built infrastructure"],
            }
        ],
        "education": [
            {
                "degree": "B.Eng Civil",
                "institution": "Tech University",
            }
        ],
        "projects": [
            {
                "title": "Tower Project",
                "technologies": ["Go"],
            }
        ],
    })

    def mock_llm(prompt: str) -> str:
        assert "Bob Builder" in prompt
        # Return response wrapped in markdown json code fences
        return f"```json\n{mock_llm_json}\n```"

    result = structure_resume("Bob Builder\nbob@builder.com", llm_callable=mock_llm)
    assert result.candidate_name == "Bob Builder"
    assert result.email == "bob@builder.com"
    assert "Go" in result.skills
    assert result.experience[0].company == "Constructions Inc"


def test_structure_resume_llm_fallback_on_error():
    """Verify structure_resume gracefully falls back to heuristics if LLM errors."""
    def failing_llm(prompt: str) -> str:
        raise RuntimeError("LLM API Timeout")

    result = structure_resume(SAMPLE_RESUME_TEXT, llm_callable=failing_llm)
    assert result.candidate_name == "Jane Doe"
    assert result.email == "jane.doe@example.com"


# =============================================================================
# 5. Orchestrator Pipeline Tests (parser.py)
# =============================================================================

def test_parse_resume_end_to_end_txt():
    """Verify full pipeline: extract -> clean -> structure on text content."""
    content = SAMPLE_RESUME_TEXT.encode("utf-8")
    result = parse_resume(content, "jane_resume.txt")

    assert isinstance(result, StructuredResume)
    assert result.candidate_name == "Jane Doe"
    assert result.email == "jane.doe@example.com"
    assert "Python" in result.skills
    assert result.raw_text is not None
    assert "Jane Doe" in result.raw_text


def test_parse_resume_end_to_end_docx():
    """Verify full pipeline on generated docx."""
    docx_bytes = create_sample_docx_bytes(
        paragraphs=[
            "Alice Wonder",
            "alice@wonderland.com",
            "555-987-6543",
            "Skills:",
            "Python, AI, ML",
            "Experience:",
            "Data Scientist at WonderCorp",
            "- Built ML pipelines",
        ]
    )
    result = parse_resume(docx_bytes, "alice.docx")
    assert isinstance(result, StructuredResume)
    assert result.candidate_name == "Alice Wonder"
    assert result.email == "alice@wonderland.com"
    assert "Python" in result.skills
