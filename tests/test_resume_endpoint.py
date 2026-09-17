import io
import json
from unittest.mock import patch

import docx
import pytest
from fastapi import status
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas
from sqlalchemy.exc import SQLAlchemyError

from app.database import Base, SessionLocal, engine, get_db
from app.main import app
from app.models import InterviewSession, Resume, SessionStatus
from app.services.resume_parser.parser import ResumeParsingError
from app.services.resume_parser.schema import ParsedResume, ResumeProject


@pytest.fixture(autouse=True)
def setup_database():
    """Create all tables before each test and drop them after."""
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)


@pytest.fixture
def client():
    """TestClient using the test database."""
    def override_get_db():
        db = SessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _make_pdf_bytes(
    name: str = "Jane Developer",
    email: str = "jane@example.com",
    skills: str = "Python, FastAPI, SQL, Docker",
    experience: str = "Software Engineer at TechCorp 2021-2024",
) -> bytes:
    """Generate in-memory PDF resume bytes."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    lines = [
        name,
        email,
        "",
        "Skills",
        skills,
        "",
        "Experience",
        experience,
        "",
        "Projects",
        "MockWise: Speech-to-Speech AI Mock Interview Platform",
        "",
        "Education",
        "B.S. Computer Science, 2020",
    ]
    y = 750
    for line in lines:
        c.drawString(72, y, line)
        y -= 20
    c.save()
    return buf.getvalue()


def _make_docx_bytes(
    name: str = "John Doe",
    skills: str = "Java, React, PostgreSQL",
) -> bytes:
    """Generate in-memory DOCX resume bytes."""
    doc = docx.Document()
    doc.add_heading(name, level=1)
    doc.add_paragraph("john.doe@example.com")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph(skills)
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Backend Engineer at CloudCorp")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.Sc. Software Engineering")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()



def test_upload_resume_pdf_success(client):
    """POST /sessions/{session_id}/resume successfully uploads and parses a PDF resume."""
    db = SessionLocal()
    session = InterviewSession(candidate_name="Jane Developer")
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    pdf_bytes = _make_pdf_bytes()
    files = {"file": ("resume.pdf", pdf_bytes, "application/pdf")}

    response = client.post(f"/sessions/{session_id}/resume", files=files)

    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["session_id"] == session_id
    assert "id" in data
    assert "skills_json" in data
    assert "projects_json" in data
    assert "experience_json" in data
    assert "education_json" in data
    assert "created_at" in data

    # Verify JSON strings are valid JSON
    skills = json.loads(data["skills_json"])
    assert isinstance(skills, list)
    assert any("Python" in s or "FastAPI" in s for s in skills)

    # Verify session fields in database
    db = SessionLocal()
    session_in_db = db.query(InterviewSession).filter_by(id=session_id).first()
    assert session_in_db.resume_filename == "resume.pdf"
    assert session_in_db.resume is not None
    assert session_in_db.resume.id == data["id"]
    db.close()


def test_upload_resume_docx_success(client):
    """POST /sessions/{session_id}/resume successfully uploads and parses a DOCX resume."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    docx_bytes = _make_docx_bytes(name="John Doe", skills="Java, React, PostgreSQL")
    files = {
        "file": (
            "candidate_cv.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    response = client.post(f"/sessions/{session_id}/resume", files=files)

    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["session_id"] == session_id
    assert "id" in data

    # Verify session filename update
    db = SessionLocal()
    session_in_db = db.query(InterviewSession).filter_by(id=session_id).first()
    assert session_in_db.resume_filename == "candidate_cv.docx"
    db.close()


def test_upload_resume_nonexistent_session_404(client):
    """POST /sessions/{session_id}/resume returns 404 when session_id does not exist."""
    pdf_bytes = _make_pdf_bytes()
    files = {"file": ("resume.pdf", pdf_bytes, "application/pdf")}

    response = client.post("/sessions/nonexistent-session-uuid/resume", files=files)

    assert response.status_code == status.HTTP_404_NOT_FOUND
    assert "not found" in response.json()["detail"].lower()


def test_upload_resume_unsupported_file_type_400(client):
    """POST /sessions/{session_id}/resume rejects non-PDF and non-DOCX files with 400."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    unsupported_files = [
        ("resume.txt", b"Plain text resume content", "text/plain"),
        ("resume.png", b"\x89PNG\r\n\x1a\nfake image", "image/png"),
        ("resume.json", b'{"skills": ["python"]}', "application/json"),
        ("resume.exe", b"MZ\x90\x00executable", "application/octet-stream"),
    ]

    for fname, content, mime in unsupported_files:
        files = {"file": (fname, content, mime)}
        response = client.post(f"/sessions/{session_id}/resume", files=files)
        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert "unsupported" in response.json()["detail"].lower() or "pdf and docx" in response.json()["detail"].lower()


def test_upload_resume_empty_file_400(client):
    """POST /sessions/{session_id}/resume rejects empty files with 400."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post(f"/sessions/{session_id}/resume", files=files)

    assert response.status_code == status.HTTP_400_BAD_REQUEST
    assert "empty" in response.json()["detail"].lower()


def test_upload_resume_parsing_error_500(client):
    """POST /sessions/{session_id}/resume returns a clean 500 when parse_resume raises ResumeParsingError."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    pdf_bytes = _make_pdf_bytes()
    files = {"file": ("corrupt.pdf", pdf_bytes, "application/pdf")}

    with patch("app.routers.sessions.parse_resume", side_effect=ResumeParsingError("Internal extraction failed")):
        response = client.post(f"/sessions/{session_id}/resume", files=files)
        assert response.status_code == status.HTTP_500_INTERNAL_SERVER_ERROR
        detail = response.json()["detail"]
        assert detail == "Failed to parse resume content."
        # Ensure internal raw trace message does not leak
        assert "Internal extraction failed" not in detail


def test_reupload_resume_overwrites_existing_record(client):
    """POST /sessions/{session_id}/resume overwrites/updates the existing Resume record for the session."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    # First upload (v1 PDF)
    pdf_bytes_v1 = _make_pdf_bytes(name="Version One", skills="C++, Go")
    res1 = client.post(
        f"/sessions/{session_id}/resume",
        files={"file": ("v1.pdf", pdf_bytes_v1, "application/pdf")},
    )
    assert res1.status_code == status.HTTP_201_CREATED
    v1_id = res1.json()["id"]

    # Second upload (v2 DOCX) for the same session
    docx_bytes_v2 = _make_docx_bytes(name="Version Two", skills="Rust, TypeScript")
    res2 = client.post(
        f"/sessions/{session_id}/resume",
        files={
            "file": (
                "v2.docx",
                docx_bytes_v2,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert res2.status_code == status.HTTP_201_CREATED
    v2_data = res2.json()

    # ID is preserved (updated in place)
    assert v2_data["id"] == v1_id
    assert v2_data["session_id"] == session_id

    # Verify DB contains only one resume record for this session
    db = SessionLocal()
    all_resumes = db.query(Resume).filter_by(session_id=session_id).all()
    assert len(all_resumes) == 1
    assert all_resumes[0].id == v1_id

    # Verify session resume_filename is updated
    updated_session = db.query(InterviewSession).filter_by(id=session_id).first()
    assert updated_session.resume_filename == "v2.docx"
    db.close()
