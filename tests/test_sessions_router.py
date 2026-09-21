import pytest
from fastapi import status
from fastapi.testclient import TestClient
from unittest.mock import patch
from sqlalchemy.exc import SQLAlchemyError

from app.database import Base, engine, get_db, SessionLocal
from app.main import app
from app.models import Answer, InterviewSession, Question, Resume, SessionStatus


@pytest.fixture(autouse=True)
def setup_database():
    """Create all tables before each test and drop them after."""
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)


@pytest.fixture
def client():
    """TestClient using the test database."""
    def override_get_db():
        db = SessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_create_session_success(client):
    """POST /sessions creates a new session and returns 201."""
    payload = {
        "candidate_name": "Jane Developer",
        "resume_filename": "jane_resume.pdf",
    }
    response = client.post("/sessions", json=payload)
    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert "id" in data
    assert data["candidate_name"] == "Jane Developer"
    assert data["resume_filename"] == "jane_resume.pdf"
    assert data["status"] == "in_progress"
    assert data["questions"] == []
    assert data["resume"] is None


def test_create_session_defaults(client):
    """POST /sessions with empty payload defaults to None."""
    response = client.post("/sessions", json={})
    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["candidate_name"] is None
    assert data["resume_filename"] is None
    assert data["status"] == "in_progress"


def test_create_session_database_error(client):
    """POST /sessions returns 500 when database error occurs."""
    with patch("app.routers.sessions.Session.commit", side_effect=SQLAlchemyError("DB down")):
        response = client.post("/sessions", json={"candidate_name": "Error Test"})
        assert response.status_code == status.HTTP_500_INTERNAL_SERVER_ERROR
        assert response.json()["detail"] == "Failed to create session due to a database error."


def test_get_session_success(client):
    """GET /sessions/{session_id} returns session with nested questions and answers."""
    # Seed session
    db = SessionLocal()
    session = InterviewSession(candidate_name="Alice", resume_filename="alice.pdf")
    db.add(session)
    db.commit()
    db.refresh(session)

    resume = Resume(session_id=session.id, raw_text="Sample resume text")
    q1 = Question(
        session_id=session.id,
        question_index=1,
        question_text="What is FastAPI?",
        question_type="technical",
    )
    db.add_all([resume, q1])
    db.commit()
    db.refresh(q1)

    a1 = Answer(question_id=q1.id, transcript_text="A modern web framework", score=9.0)
    db.add(a1)
    db.commit()
    session_id = session.id
    db.close()

    response = client.get(f"/sessions/{session_id}")
    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data["id"] == session.id
    assert data["candidate_name"] == "Alice"
    assert data["resume"] is not None
    assert data["resume"]["raw_text"] == "Sample resume text"
    assert len(data["questions"]) == 1
    assert data["questions"][0]["question_text"] == "What is FastAPI?"
    assert data["questions"][0]["answer"]["transcript_text"] == "A modern web framework"
    assert data["questions"][0]["answer"]["score"] == 9.0


def test_get_session_not_found(client):
    """GET /sessions/{session_id} returns 404 for non-existent session."""
    response = client.get("/sessions/non-existent-id")
    assert response.status_code == status.HTTP_404_NOT_FOUND
    assert "not found" in response.json()["detail"].lower()


def test_get_session_database_error(client):
    """GET /sessions/{session_id} returns 500 when database query fails."""
    with patch("app.routers.sessions.Session.query", side_effect=SQLAlchemyError("DB failure")):
        response = client.get("/sessions/some-id")
        assert response.status_code == status.HTTP_500_INTERNAL_SERVER_ERROR
        assert response.json()["detail"] == "Failed to retrieve session due to a database error."


def test_complete_session_success(client):
    """POST /sessions/{session_id}/complete marks session completed."""
    db = SessionLocal()
    session = InterviewSession(candidate_name="Bob")
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    response = client.post(f"/sessions/{session_id}/complete")
    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data["id"] == session_id
    assert data["status"] == "completed"

    # Verify persistent state in DB
    db = SessionLocal()
    updated = db.query(InterviewSession).filter_by(id=session_id).first()
    assert updated.status == SessionStatus.completed
    db.close()


def test_complete_session_not_found(client):
    """POST /sessions/{session_id}/complete returns 404 when session does not exist."""
    response = client.post("/sessions/non-existent-id/complete")
    assert response.status_code == status.HTTP_404_NOT_FOUND
    assert "not found" in response.json()["detail"].lower()


def test_complete_session_database_error_on_commit(client):
    """POST /sessions/{session_id}/complete returns 500 and rolls back when commit fails."""
    db = SessionLocal()
    session = InterviewSession(candidate_name="Charlie")
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    with patch("app.routers.sessions.Session.commit", side_effect=SQLAlchemyError("Commit failed")):
        response = client.post(f"/sessions/{session_id}/complete")
        assert response.status_code == status.HTTP_500_INTERNAL_SERVER_ERROR
        assert response.json()["detail"] == "Failed to update session due to a database error."


# =============================================================================
# Resume Upload & Parsing Endpoint Tests (POST /sessions/{session_id}/resume)
# =============================================================================

def _make_sample_pdf_bytes() -> bytes:
    import io
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    lines = [
        "Sarah Connor",
        "sarah@skynet.resistance",
        "",
        "Skills",
        "Tactics, Python, Security",
        "",
        "Experience",
        "Leader at Resistance - Directed tactical operations",
    ]
    y = 750
    for line in lines:
        c.drawString(72, y, line)
        y -= 20
    c.save()
    return buf.getvalue()



def _make_sample_docx_bytes() -> bytes:
    import io
    import docx
    doc = docx.Document()
    doc.add_paragraph("John Doe")
    doc.add_paragraph("john@example.com")
    doc.add_paragraph("Skills:\nPython, Machine Learning")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_resume_success_pdf(client):
    """POST /sessions/{session_id}/resume successfully processes a PDF resume."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    pdf_bytes = _make_sample_pdf_bytes()
    files = {"file": ("sarah_resume.pdf", pdf_bytes, "application/pdf")}
    response = client.post(f"/sessions/{session_id}/resume", files=files)

    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["session_id"] == session_id
    assert "id" in data
    assert "skills_json" in data
    assert "Python" in data["skills_json"]

    # Verify session fields updated
    db = SessionLocal()
    updated_session = db.query(InterviewSession).filter_by(id=session_id).first()
    assert updated_session.resume_filename == "sarah_resume.pdf"
    assert updated_session.resume is not None
    db.close()

    # Verify GET /sessions/{session_id} returns nested resume
    get_res = client.get(f"/sessions/{session_id}")
    assert get_res.status_code == status.HTTP_200_OK
    assert get_res.json()["resume"]["id"] == data["id"]


def test_upload_resume_success_docx(client):
    """POST /sessions/{session_id}/resume successfully processes a DOCX resume."""
    docx_bytes = _make_sample_docx_bytes()

    db = SessionLocal()
    session = InterviewSession(candidate_name="Existing Candidate")
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    files = {
        "file": (
            "john.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post(f"/sessions/{session_id}/resume", files=files)

    assert response.status_code == status.HTTP_201_CREATED
    data = response.json()
    assert data["session_id"] == session_id
    assert "Python" in data["skills_json"]

    # Candidate name should not be overwritten if already set
    db = SessionLocal()
    updated_session = db.query(InterviewSession).filter_by(id=session_id).first()
    assert updated_session.candidate_name == "Existing Candidate"
    assert updated_session.resume_filename == "john.docx"
    db.close()


def test_upload_resume_session_not_found(client):
    """POST /sessions/{session_id}/resume returns 404 for non-existent session."""
    pdf_bytes = _make_sample_pdf_bytes()
    files = {"file": ("test.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/sessions/unknown-session-id/resume", files=files)
    assert response.status_code == status.HTTP_404_NOT_FOUND
    assert "not found" in response.json()["detail"].lower()


def test_upload_resume_empty_file(client):
    """POST /sessions/{session_id}/resume returns 400 when file is empty."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post(f"/sessions/{session_id}/resume", files=files)
    assert response.status_code == status.HTTP_400_BAD_REQUEST
    assert "empty" in response.json()["detail"].lower()


def test_upload_resume_unsupported_type(client):
    """POST /sessions/{session_id}/resume returns 400 when file type is unsupported."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    files = {"file": ("malware.exe", b"binarycontent", "application/octet-stream")}
    response = client.post(f"/sessions/{session_id}/resume", files=files)
    assert response.status_code == status.HTTP_400_BAD_REQUEST
    assert "unsupported file format" in response.json()["detail"].lower()


def test_upload_resume_updates_existing_resume(client):
    """POST /sessions/{session_id}/resume replaces/updates existing 1:1 resume."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    pdf_bytes1 = _make_sample_pdf_bytes()
    docx_bytes2 = _make_sample_docx_bytes()

    # First upload
    res1 = client.post(
        f"/sessions/{session_id}/resume",
        files={"file": ("v1.pdf", pdf_bytes1, "application/pdf")},
    )
    assert res1.status_code == status.HTTP_201_CREATED
    first_id = res1.json()["id"]

    # Second upload for same session
    res2 = client.post(
        f"/sessions/{session_id}/resume",
        files={
            "file": (
                "v2.docx",
                docx_bytes2,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert res2.status_code == status.HTTP_201_CREATED
    second_data = res2.json()

    # Id should be preserved as the existing 1:1 record is updated
    assert second_data["id"] == first_id
    assert "skills_json" in second_data

    # Ensure only 1 Resume record exists in DB for this session
    db = SessionLocal()
    resumes = db.query(Resume).filter_by(session_id=session_id).all()
    assert len(resumes) == 1
    db.close()


def test_upload_resume_database_error(client):
    """POST /sessions/{session_id}/resume returns 500 when database save fails."""
    db = SessionLocal()
    session = InterviewSession()
    db.add(session)
    db.commit()
    session_id = session.id
    db.close()

    pdf_bytes = _make_sample_pdf_bytes()
    with patch("app.routers.sessions.Session.commit", side_effect=SQLAlchemyError("DB save failed")):
        response = client.post(
            f"/sessions/{session_id}/resume",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
        )
        assert response.status_code == status.HTTP_500_INTERNAL_SERVER_ERROR
        assert response.json()["detail"] == "Failed to save resume due to a database error."


def test_complete_session_ignores_skipped_answers(client):
    """POST /sessions/{session_id}/complete leaves skipped answers unscored without error."""
    db = SessionLocal()
    session = InterviewSession(candidate_name="Skipper")
    db.add(session)
    db.commit()

    q1 = Question(session_id=session.id, question_index=0, question_text="Q1 text")
    q2 = Question(session_id=session.id, question_index=1, question_text="Q2 text")
    db.add_all([q1, q2])
    db.commit()

    # q1 is answered with transcript, unscored
    a1 = Answer(question_id=q1.id, transcript_text="Detailed answer to question 1", skipped=False)
    # q2 is skipped
    a2 = Answer(question_id=q2.id, transcript_text=None, score=None, skipped=True)
    db.add_all([a1, a2])
    db.commit()
    session_id = session.id
    q1_id, q2_id = q1.id, q2.id
    db.close()

    res = client.post(f"/sessions/{session_id}/complete")
    assert res.status_code == status.HTTP_200_OK
    data = res.json()
    assert data["status"] == "completed"
    assert data["answered_count"] == 1
    assert data["skipped_count"] == 1
    assert data["total_questions"] == 2

    # Check question answers
    q_map = {q["id"]: q for q in data["questions"]}
    assert q_map[q1_id]["answer"]["score"] is not None
    assert q_map[q1_id]["answer"]["skipped"] is False
    assert q_map[q2_id]["answer"]["score"] is None
    assert q_map[q2_id]["answer"]["skipped"] is True



def test_session_summary_mixed_answered_skipped_untouched(client):
    """GET /sessions/{session_id} computes correct answered_count, skipped_count, and average_score."""
    db = SessionLocal()
    session = InterviewSession(candidate_name="Mixed Candidate")
    db.add(session)
    db.commit()

    q0 = Question(session_id=session.id, question_index=0, question_text="Q0", question_type="resume_specific")
    q1 = Question(session_id=session.id, question_index=1, question_text="Q1", question_type="general")
    q2 = Question(session_id=session.id, question_index=2, question_text="Q2", question_type="general")
    q3 = Question(session_id=session.id, question_index=3, question_text="Q3", question_type="general")
    db.add_all([q0, q1, q2, q3])
    db.commit()

    # q0 answered with score 8.0
    a0 = Answer(question_id=q0.id, transcript_text="Ans 0", score=8.0, skipped=False)
    # q1 answered with score 6.0
    a1 = Answer(question_id=q1.id, transcript_text="Ans 1", score=6.0, skipped=False)
    # q2 skipped
    a2 = Answer(question_id=q2.id, transcript_text=None, score=None, skipped=True)
    # q3 untouched (no Answer row)
    db.add_all([a0, a1, a2])
    db.commit()
    session_id = session.id
    db.close()

    res = client.get(f"/sessions/{session_id}")
    assert res.status_code == status.HTTP_200_OK
    data = res.json()
    assert data["total_questions"] == 4
    assert data["answered_count"] == 2
    assert data["skipped_count"] == 1
    assert data["average_score"] == 7.0  # (8.0 + 6.0) / 2
    assert data["resume_specific_count"] == 1
    assert data["general_count"] == 3


