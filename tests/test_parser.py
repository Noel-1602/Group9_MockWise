"""Tests for parse_resume orchestrator in app/services/resume_parser/parser.py.

Tests end-to-end parsing of PDF and DOCX files using the "mock" backend,
unsupported file format handling, and validation retry logic.
"""

import pathlib
from unittest.mock import patch

import pytest

from app.services.resume_parser.extract import UnsupportedFileTypeError
from app.services.resume_parser.parser import ResumeParsingError, parse_resume
from app.services.resume_parser.schema import ParsedResume, ResumeProject


# ---------------------------------------------------------------------------
# Helpers -- build sample files with realistic resume-like content
# ---------------------------------------------------------------------------

SAMPLE_NAME = "Jane Doe"
SAMPLE_EMAIL = "jane.doe@example.com"
SAMPLE_SKILLS = "Python, FastAPI, PostgreSQL, Docker"
SAMPLE_PROJECT = "MockWise: an AI-powered mock interview platform built with FastAPI"
SAMPLE_EDUCATION = "B.Sc. Computer Science, University of Technology, 2022"


def _build_sample_pdf(dest: pathlib.Path) -> None:
    """Write a minimal, text-searchable PDF using reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(dest), pagesize=A4)
    _, height = A4

    lines_to_write = [
        SAMPLE_NAME,
        SAMPLE_EMAIL,
        "",
        "Skills",
        SAMPLE_SKILLS,
        "",
        "Projects",
        SAMPLE_PROJECT,
        "",
        "Education",
        SAMPLE_EDUCATION,
    ]

    y = height - 72  # start 1 inch from the top
    for line in lines_to_write:
        c.drawString(72, y, line)
        y -= 18  # 18-point line spacing

    c.save()


def _build_sample_docx(dest: pathlib.Path) -> None:
    """Write a minimal DOCX with paragraphs using python-docx."""
    import docx

    doc = docx.Document()
    doc.add_heading(SAMPLE_NAME, level=1)
    doc.add_paragraph(SAMPLE_EMAIL)
    doc.add_heading("Skills", level=2)
    doc.add_paragraph(SAMPLE_SKILLS)
    doc.add_heading("Projects", level=2)
    doc.add_paragraph(SAMPLE_PROJECT)
    doc.add_heading("Education", level=2)
    doc.add_paragraph(SAMPLE_EDUCATION)
    doc.save(str(dest))


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture()
def sample_pdf(tmp_path: pathlib.Path) -> pathlib.Path:
    """Return the path to a temporary sample PDF."""
    pdf_path = tmp_path / "sample_resume.pdf"
    _build_sample_pdf(pdf_path)
    return pdf_path


@pytest.fixture()
def sample_docx(tmp_path: pathlib.Path) -> pathlib.Path:
    """Return the path to a temporary sample DOCX."""
    docx_path = tmp_path / "sample_resume.docx"
    _build_sample_docx(docx_path)
    return docx_path


# ---------------------------------------------------------------------------
# Tests: PDF End-to-End
# ---------------------------------------------------------------------------

class TestParseResumePdf:
    """End-to-end tests for parse_resume with PDF files."""

    def test_returns_parsed_resume_instance(self, sample_pdf: pathlib.Path) -> None:
        result = parse_resume(str(sample_pdf), backend="mock")
        assert isinstance(result, ParsedResume)

    def test_extracts_skills(self, sample_pdf: pathlib.Path) -> None:
        result = parse_resume(str(sample_pdf), backend="mock")
        assert len(result.skills) > 0
        skill_lower = [s.lower() for s in result.skills]
        assert "python" in skill_lower or "fastapi" in skill_lower

    def test_extracts_projects(self, sample_pdf: pathlib.Path) -> None:
        result = parse_resume(str(sample_pdf), backend="mock")
        assert len(result.projects) > 0
        assert any("mockwise" in p.title.lower() or "mockwise" in p.description.lower() for p in result.projects)

    def test_extracts_education(self, sample_pdf: pathlib.Path) -> None:
        result = parse_resume(str(sample_pdf), backend="mock")
        assert len(result.education) > 0
        combined = " ".join(result.education).lower()
        assert "computer science" in combined or "technology" in combined


# ---------------------------------------------------------------------------
# Tests: DOCX End-to-End
# ---------------------------------------------------------------------------

class TestParseResumeDocx:
    """End-to-end tests for parse_resume with DOCX files."""

    def test_returns_parsed_resume_instance(self, sample_docx: pathlib.Path) -> None:
        result = parse_resume(str(sample_docx), backend="mock")
        assert isinstance(result, ParsedResume)

    def test_extracts_skills(self, sample_docx: pathlib.Path) -> None:
        result = parse_resume(str(sample_docx), backend="mock")
        assert len(result.skills) > 0
        skill_lower = [s.lower() for s in result.skills]
        assert "python" in skill_lower or "docker" in skill_lower

    def test_extracts_projects(self, sample_docx: pathlib.Path) -> None:
        result = parse_resume(str(sample_docx), backend="mock")
        assert len(result.projects) > 0
        assert any("mockwise" in p.title.lower() for p in result.projects)

    def test_extracts_education(self, sample_docx: pathlib.Path) -> None:
        result = parse_resume(str(sample_docx), backend="mock")
        assert len(result.education) > 0
        combined = " ".join(result.education).lower()
        assert "computer science" in combined or "university" in combined


# ---------------------------------------------------------------------------
# Tests: Default Backend
# ---------------------------------------------------------------------------

class TestParseResumeDefaultBackend:
    """Test that default backend is "mock"."""

    def test_default_backend_is_mock(self, sample_pdf: pathlib.Path) -> None:
        result_default = parse_resume(str(sample_pdf))
        result_explicit = parse_resume(str(sample_pdf), backend="mock")
        assert isinstance(result_default, ParsedResume)
        assert result_default.skills == result_explicit.skills
        assert len(result_default.projects) == len(result_explicit.projects)
        assert result_default.education == result_explicit.education


# ---------------------------------------------------------------------------
# Tests: Unsupported File Extensions
# ---------------------------------------------------------------------------

class TestParseResumeUnsupportedFormats:
    """Test that unsupported file formats raise clear exceptions."""

    def test_txt_raises_unsupported_error(self, tmp_path: pathlib.Path) -> None:
        txt_path = tmp_path / "resume.txt"
        txt_path.write_text("Jane Doe\nSkills\nPython, FastAPI")
        with pytest.raises((UnsupportedFileTypeError, ResumeParsingError)):
            parse_resume(str(txt_path))

    def test_other_unsupported_extension_raises(self, tmp_path: pathlib.Path) -> None:
        bad_path = tmp_path / "resume.png"
        bad_path.write_bytes(b"PNG fake binary content")
        with pytest.raises((UnsupportedFileTypeError, ResumeParsingError)):
            parse_resume(str(bad_path))


# ---------------------------------------------------------------------------
# Tests: Validation & Retry Logic
# ---------------------------------------------------------------------------

class TestParseResumeRetryLogic:
    """Test retry behavior when structure_resume produces invalid output."""

    def test_retry_succeeds_after_transient_validation_failure(
        self, sample_pdf: pathlib.Path
    ) -> None:
        """If first structuring attempt returns malformed data, retry recovers."""
        valid_dict = {
            "skills": ["Python", "FastAPI"],
            "projects": [{"title": "TestProj", "description": "Desc"}],
            "experience": ["Engineer"],
            "education": ["B.Sc."],
        }
        calls = [{"skills": 12345}, valid_dict]

        def mock_structure(*args, **kwargs):
            return calls.pop(0)

        with patch("app.services.resume_parser.parser.structure_resume", side_effect=mock_structure):
            result = parse_resume(str(sample_pdf), backend="mock")
            assert isinstance(result, ParsedResume)
            assert result.skills == ["Python", "FastAPI"]

    def test_retry_failure_raises_resume_parsing_error(
        self, sample_pdf: pathlib.Path
    ) -> None:
        """If structuring/validation fails twice, ResumeParsingError is raised."""
        malformed_dict = {"skills": 12345}

        with patch("app.services.resume_parser.parser.structure_resume", return_value=malformed_dict):
            with pytest.raises(ResumeParsingError, match="Failed to structure and validate resume"):
                parse_resume(str(sample_pdf), backend="mock")
