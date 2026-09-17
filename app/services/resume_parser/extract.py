"""Raw text extraction from various resume document formats (PDF, DOCX, TXT)."""

import io
import os
from typing import Set

import docx
import pypdf


class ResumeExtractionError(Exception):
    """Base exception for resume text extraction failures."""
    pass


class UnsupportedFileTypeError(ResumeExtractionError):
    """Raised when an uploaded file type is not supported."""
    pass


class CorruptedFileError(ResumeExtractionError):
    """Raised when file content cannot be parsed or is corrupted."""
    pass


class EmptyFileError(ResumeExtractionError):
    """Raised when the uploaded file contains zero bytes."""
    pass


SUPPORTED_EXTENSIONS: Set[str] = {".pdf", ".docx", ".txt", ".md"}


def extract_from_pdf(content: bytes) -> str:
    """Extract text from PDF binary content using pypdf."""
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages_text = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                pages_text.append(page_text)
        return "\n\n".join(pages_text)
    except Exception as exc:
        raise CorruptedFileError(f"Failed to extract text from PDF: {exc}") from exc


def extract_from_docx(content: bytes) -> str:
    """Extract text from DOCX binary content using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also include any table text
        table_cells = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_cells.append(row_text)

        all_text = paragraphs + table_cells
        return "\n".join(all_text)
    except Exception as exc:
        raise CorruptedFileError(f"Failed to extract text from DOCX: {exc}") from exc


def extract_from_txt(content: bytes) -> str:
    """Extract text from plain text or markdown binary content with encoding fallbacks."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise CorruptedFileError("Could not decode text file with supported encodings (UTF-8, Latin-1).")


def extract_text(content: bytes, filename: str) -> str:
    """Extract raw text from a document based on its file extension.

    Args:
        content: Raw binary content of the file.
        filename: Name of the file including extension.

    Returns:
        Extracted raw string content.

    Raises:
        EmptyFileError: If content is empty.
        UnsupportedFileTypeError: If file extension is unsupported.
        CorruptedFileError: If file cannot be parsed.
    """
    if not content or len(content) == 0:
        raise EmptyFileError(f"File '{filename}' is empty.")

    ext = os.path.splitext(filename)[1].lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file format '{ext}'. Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        return extract_from_pdf(content)
    elif ext == ".docx":
        return extract_from_docx(content)
    elif ext in (".txt", ".md"):
        return extract_from_txt(content)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file extension: {ext}")


# ---------------------------------------------------------------------------
# Path-based extraction functions (used by the resume-parser pipeline)
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file_path: str) -> str:
    """Extract raw text from a PDF file using PyMuPDF (pymupdf).

    No cleaning or post-processing is applied; that is handled downstream.
    If PyMuPDF returns empty or very short text the raw result is returned
    as-is — no fallback is attempted here.

    Args:
        file_path: Absolute or relative path to the PDF file.

    Returns:
        Raw text concatenated from every page, separated by newlines.
    """
    import pymupdf  # PyMuPDF – imported locally to keep the module importable
                    # even when the optional dependency is absent.

    pages: list[str] = []
    with pymupdf.open(file_path) as doc:
        for page in doc:
            pages.append(page.get_text())
    return "\n".join(pages)


def extract_text_from_docx(file_path: str) -> str:
    """Extract raw text from a DOCX file using python-docx.

    Paragraph text and table cell text are both included.  No cleaning or
    post-processing is applied; that is handled downstream.

    Args:
        file_path: Absolute or relative path to the DOCX file.

    Returns:
        Raw text with one paragraph/row per line.
    """
    import docx as _docx  # aliased to avoid shadowing the module-level import

    doc = _docx.Document(file_path)
    lines: list[str] = []

    for para in doc.paragraphs:
        lines.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))

    return "\n".join(lines)
